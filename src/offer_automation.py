from __future__ import annotations

import argparse
import base64
import json
import logging
import re
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urljoin, urlparse

import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


LOGGER = logging.getLogger("offer_automation")
SEARCH_URL = "https://etilbudsavis.dk/soeg/{}"
CATEGORY_ORDER = ("Personal care", "Home care")


@dataclass(frozen=True)
class Offer:
    name: str
    url: str
    image_url: str
    retailer: str
    price: str
    valid_through: str
    category: str = ""
    publication_id: str = ""
    description: str = ""


@dataclass(frozen=True)
class Publication:
    id: str
    label: str
    cover_image: str
    valid_from: str
    valid_until: str
    published_at: str
    source: str = "etilbudsavis"
    page_url: str = ""
    api_key: str = field(default="", repr=False)
    dealer_id: str = ""
    offer_count: int = 0


class JsonLdParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._is_json_ld = False
        self._script_parts: list[str] = []
        self.scripts: list[str] = []
        self._is_app_data = False
        self._app_data_key = ""
        self._app_data_parts: list[str] = []
        self.app_data: list[tuple[str, str]] = []
        self.coop_api_key = ""
        self.incito_publication_id = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag.lower() == "app-data":
            self._is_app_data = True
            self._app_data_key = attributes.get("data-key") or ""
            self._app_data_parts = []
            return
        if tag.lower() == "div" and "incito-wrapper" in (attributes.get("class") or "").split():
            self.incito_publication_id = attributes.get("data-publication-id") or ""
        if tag.lower() == "script" and attributes.get("id") == "sgn-sdk":
            self.coop_api_key = attributes.get("data-app-key") or ""
        if tag.lower() != "script":
            return
        self._is_json_ld = (attributes.get("type") or "").lower() == "application/ld+json"
        self._script_parts = []

    def handle_data(self, data: str) -> None:
        if self._is_json_ld:
            self._script_parts.append(data)
        if self._is_app_data:
            self._app_data_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "app-data" and self._is_app_data:
            self.app_data.append((self._app_data_key, "".join(self._app_data_parts)))
            self._is_app_data = False
            self._app_data_key = ""
            self._app_data_parts = []
        if tag.lower() == "script" and self._is_json_ld:
            self.scripts.append("".join(self._script_parts))
            self._is_json_ld = False
            self._script_parts = []


class CoopCoverParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._in_slider = False
        self._slider_depth = 0
        self._current_url = ""
        self.covers: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag.lower() == "div" and not self._in_slider:
            classes = (attributes.get("class") or "").split()
            if "catalog-images-slider" in classes and "coop-slider" in classes:
                self._in_slider = True
                self._slider_depth = 1
                return
        if not self._in_slider:
            return
        if tag.lower() == "div":
            self._slider_depth += 1
        elif tag.lower() == "a":
            self._current_url = attributes.get("href") or ""
        elif tag.lower() == "img" and self._current_url:
            image_url = attributes.get("src") or attributes.get("data-src") or ""
            if image_url:
                self.covers.append((self._current_url, image_url))

    def handle_endtag(self, tag: str) -> None:
        if not self._in_slider:
            return
        if tag.lower() == "a":
            self._current_url = ""
        elif tag.lower() == "div":
            self._slider_depth -= 1
            if self._slider_depth == 0:
                self._in_slider = False


def _walk_json(value: Any):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_json(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_json(child)


def parse_search_results(page_html: str) -> list[Offer]:
    parser = JsonLdParser()
    parser.feed(page_html)
    offers: list[Offer] = []
    seen_urls: set[str] = set()

    for script in parser.scripts:
        try:
            structured_data = json.loads(script)
        except json.JSONDecodeError:
            LOGGER.warning("Ignorerer ugyldig JSON-LD fra søgeresultatet")
            continue

        for record in _walk_json(structured_data):
            record_type = record.get("@type", [])
            if isinstance(record_type, str):
                record_type = [record_type]
            if "Offer" not in record_type:
                continue

            seller = record.get("seller", {})
            if isinstance(seller, list):
                seller = seller[0] if seller else {}
            name = record.get("name")
            url = record.get("url")
            image_url = record.get("image")
            if isinstance(image_url, list):
                image_url = image_url[0] if image_url else ""
            retailer = seller.get("name", "") if isinstance(seller, dict) else str(seller)
            if not all(isinstance(value, str) and value for value in (name, url, image_url, retailer)):
                continue
            if url in seen_urls:
                continue
            seen_urls.add(url)
            publication_id = parse_qs(urlparse(url).query).get("publication", [""])[0]
            offers.append(
                Offer(
                    name=name.strip(),
                    url=url,
                    image_url=image_url,
                    retailer=retailer.strip(),
                    price=str(record.get("price", "")),
                    valid_through=str(record.get("validThrough", record.get("priceValidUntil", ""))),
                    publication_id=publication_id,
                    description=str(record.get("description", "")),
                )
            )

    for encoded_key, payload in parser.app_data:
        try:
            data_key = json.loads(base64.b64decode(encoded_key).decode("utf-8"))
            if (not data_key or data_key[0] != "offers"
                    or data_key[1].get("sources") != ["publication"]):
                continue
            result = json.loads(payload)
        except (ValueError, TypeError, UnicodeDecodeError, json.JSONDecodeError):
            continue

        for record in result.get("data", []):
            business = record.get("business", {})
            business_slugs = business.get("slugs", []) if isinstance(business, dict) else []
            business_name = business.get("name", "") if isinstance(business, dict) else ""
            offer_id = record.get("publicId", "")
            publication_id = record.get("publicationPublicId", "")
            image_url = record.get("imageLarge") or record.get("image")
            if not all((offer_id, publication_id, image_url, record.get("name"), business_name)):
                continue
            if not business_slugs:
                continue
            offer_url = (
                f"https://etilbudsavis.dk/{quote(business_slugs[0], safe='')}"
                f"?publication={quote(publication_id, safe='')}&offer={quote(offer_id, safe='')}"
            )
            if offer_url in seen_urls:
                continue
            seen_urls.add(offer_url)
            offers.append(
                Offer(
                    name=record["name"].strip(),
                    url=offer_url,
                    image_url=image_url,
                    retailer=business_name.strip(),
                    price=str(record.get("price", "")),
                    valid_through=str(record.get("validUntil", "")),
                    publication_id=publication_id,
                    description=str(record.get("description", "")),
                )
            )
    return offers


def parse_publications(page_html: str) -> list[Publication]:
    parser = JsonLdParser()
    parser.feed(page_html)
    publications: list[Publication] = []

    for encoded_key, payload in parser.app_data:
        try:
            key = json.loads(base64.b64decode(encoded_key).decode("utf-8"))
            if not key or key[0] != "fronts":
                continue
            groups = json.loads(payload)
        except (ValueError, UnicodeDecodeError, json.JSONDecodeError):
            continue

        for group in groups:
            for item in group.get("publications", []):
                publication_id = item.get("id")
                cover_image = item.get("view") or item.get("thumb")
                if not publication_id or not cover_image:
                    continue
                publications.append(
                    Publication(
                        id=publication_id,
                        label=item.get("label", "Tilbudsavis"),
                        cover_image=cover_image,
                        valid_from=item.get("validFrom", ""),
                        valid_until=item.get("validUntil", ""),
                        published_at=item.get("publish", ""),
                    )
                )
    return publications


def select_week_publication(publications: list[Publication], week_number: int) -> Publication | None:
    pattern = re.compile(r"\b(?:uge|week)\s*\(?([0-9]{1,2})\)?\b", re.IGNORECASE)
    matches = [
        publication
        for publication in publications
        if (match := pattern.search(publication.label)) and int(match.group(1)) == week_number
    ]
    if not matches:
        return None
    return max(matches, key=lambda publication: publication.published_at or publication.valid_from)


def _normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.casefold())
    return "".join(char for char in normalized if not unicodedata.combining(char))


def _term_matches(text: str, term: str) -> bool:
    normalized_term = _normalize(term).strip()
    normalized_text = _normalize(text)
    return bool(re.search(r"(?<!\w)" + re.escape(normalized_term) + r"(?!\w)", normalized_text))


def categorize_offer(name: str, categories: dict[str, dict[str, list[str]]],
                     excluded_terms: list[str] | None = None,
                     description: str = "") -> str | None:
    searchable_text = f"{name} {description}"
    if any(_term_matches(searchable_text, term) for term in (excluded_terms or [])):
        return None
    matched: list[tuple[int, str]] = []
    for category, groups in categories.items():
        terms = groups.get("keywords", []) + groups.get("brands", [])
        lengths = [len(term) for term in terms if _term_matches(searchable_text, term)]
        if lengths:
            matched.append((max(lengths), category))
    if not matched:
        return None
    best_score = max(score for score, _ in matched)
    winners = [category for score, category in matched if score == best_score]
    return winners[0] if len(winners) == 1 else None


def _retailer_display_name(raw_name: str, retailers: list[dict[str, Any]]) -> str | None:
    normalized = _normalize(raw_name).strip()
    for retailer in retailers:
        aliases = [retailer["name"], *retailer.get("aliases", [])]
        if normalized in {_normalize(alias).strip() for alias in aliases}:
            return retailer["name"]
    return None


def _is_current(offer: Offer, now: datetime) -> bool:
    if not offer.valid_through:
        return True
    try:
        expiry = datetime.fromisoformat(offer.valid_through.replace("Z", "+00:00"))
    except ValueError:
        return True
    if expiry.tzinfo is None:
        expiry = expiry.replace(tzinfo=timezone.utc)
    return expiry >= now


def _parse_datetime(value: str) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed


def _overlaps_week(valid_from: str, valid_until: str, week_number: int) -> bool:
    week_year = date.today().isocalendar().year
    week_start_date = date.fromisocalendar(week_year, week_number, 1)
    week_start = datetime.combine(week_start_date, datetime.min.time(), timezone.utc)
    week_end = week_start + timedelta(days=7)
    offer_start = _parse_datetime(valid_from)
    offer_end = _parse_datetime(valid_until)
    return ((offer_start is None or offer_start < week_end)
            and (offer_end is None or offer_end >= week_start))


def _normalize_page_url(url: str) -> str:
    parsed = urlparse(url)
    return f"{parsed.scheme.lower()}://{parsed.netloc.lower()}{parsed.path.rstrip('/')}/"


def _fetch_coop_cover_images(session: requests.Session) -> dict[str, str]:
    response = session.get("https://coop.dk/butikker/ugens-aviser/", timeout=30)
    response.raise_for_status()
    parser = CoopCoverParser()
    parser.feed(response.text)
    return {
        _normalize_page_url(urljoin(response.url, page_url)): urljoin(response.url, image_url)
        for page_url, image_url in parser.covers
    }


def _fetch_coop_publication(retailer: dict[str, Any], week_number: int,
                            session: requests.Session,
                            cover_images: dict[str, str]) -> Publication | None:
    page_url = retailer["coop_page_url"]
    response = session.get(page_url, timeout=30)
    response.raise_for_status()
    parser = JsonLdParser()
    parser.feed(response.text)
    if not parser.incito_publication_id or not parser.coop_api_key:
        LOGGER.warning("Coop-avis for '%s' mangler publikations-ID eller læseadgang",
                       retailer["name"])
        return None

    headers = {"X-Api-Key": parser.coop_api_key}
    catalog_url = f"https://squid-api.tjek.com/v2/catalogs/{quote(parser.incito_publication_id, safe='')}"
    catalog_response = session.get(catalog_url, headers=headers, timeout=30)
    catalog_response.raise_for_status()
    catalog = catalog_response.json()
    if not _overlaps_week(catalog.get("run_from", ""), catalog.get("run_till", ""), week_number):
        LOGGER.warning("Coop-avis for '%s' gælder ikke uge %d", retailer["name"], week_number)
        return None

    cover_image = cover_images.get(_normalize_page_url(page_url), "")
    if not cover_image:
        LOGGER.warning("Coop-forside ikke fundet for '%s'", retailer["name"])
    return Publication(
        id=catalog["id"],
        label=catalog.get("label", f"Coop avis uge {week_number}"),
        cover_image=cover_image,
        valid_from=catalog.get("run_from", ""),
        valid_until=catalog.get("run_till", ""),
        published_at=catalog.get("publish", ""),
        source="coop",
        page_url=page_url,
        api_key=parser.coop_api_key,
        dealer_id=catalog.get("dealer_id", ""),
        offer_count=int(catalog.get("offer_count") or 0),
    )


def _fetch_coop_offers(retailer: dict[str, Any], publication: Publication,
                       week_number: int, session: requests.Session,
                       delay_seconds: float) -> list[Offer]:
    headers = {"X-Api-Key": publication.api_key}
    offers: list[Offer] = []
    offset = 0
    page_size = 100

    while True:
        response = session.get(
            "https://squid-api.tjek.com/v2/offers",
            headers=headers,
            params={
                "dealer_id": publication.dealer_id,
                "catalog_id": publication.id,
                "types": "offer",
                "order_by": "page",
                "offset": offset,
                "limit": page_size,
            },
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        records = payload if isinstance(payload, list) else payload.get("data", [])
        if not records:
            break

        for record in records:
            if not _overlaps_week(record.get("run_from", ""),
                                  record.get("run_till", ""), week_number):
                continue
            images = record.get("images", {})
            image_url = images.get("zoom") or images.get("view") or images.get("thumb")
            if not image_url:
                continue
            catalog_page = int(record.get("catalog_page") or 0) + 1
            offer_id = record.get("catalog_view_id") or record.get("id", "")
            separator = "&" if "?" in publication.page_url else "?"
            offer_url = (f"{publication.page_url}{separator}page={catalog_page}"
                         f"&offer={quote(str(offer_id), safe='')}")
            pricing = record.get("pricing", {})
            price = pricing.get("price")
            offers.append(
                Offer(
                    name=str(record.get("heading", "")).strip(),
                    url=offer_url,
                    image_url=image_url,
                    retailer=retailer["name"],
                    price=str(price) if price is not None else "",
                    valid_through=str(record.get("run_till", "")),
                    publication_id=publication.id,
                    description=str(record.get("description", "")),
                )
            )
        offset += len(records)
        if len(records) < page_size or (publication.offer_count and offset >= publication.offer_count):
            break
        time.sleep(delay_seconds)
    return offers


def _search_terms(config: dict[str, Any]) -> list[str]:
    terms = list(config.get("search_terms", []))
    for groups in config["categories"].values():
        terms.extend(groups.get("keywords", []))
        terms.extend(groups.get("brands", []))
    return list(dict.fromkeys(term.strip() for term in terms if term.strip()))


def fetch_publications(retailers: list[dict[str, Any]], week_number: int,
                       delay_seconds: float = 1.0) -> dict[str, Publication | None]:
    session = requests.Session()
    session.headers.update({"User-Agent": "OfferAutomation/1.0 (public offer search)"})
    selected: dict[str, Publication | None] = {}
    successful_pages = 0
    coop_cover_images: dict[str, str] = {}
    if any(retailer.get("coop_page_url") for retailer in retailers):
        try:
            coop_cover_images = _fetch_coop_cover_images(session)
        except requests.RequestException as exc:
            LOGGER.warning("Coops avisforsider kunne ikke hentes: %s", exc)

    for retailer in retailers:
        if retailer.get("coop_page_url"):
            try:
                publication = _fetch_coop_publication(
                    retailer, week_number, session, coop_cover_images
                )
                selected[retailer["name"]] = publication
                if publication:
                    successful_pages += 1
                    LOGGER.info("%s: valgt Coop-avis '%s' (%d tilbud)",
                                retailer["name"], publication.label, publication.offer_count)
                else:
                    LOGGER.warning("%s: ingen Coop-avis fundet for uge %d",
                                   retailer["name"], week_number)
            except (requests.RequestException, ValueError, KeyError) as exc:
                LOGGER.warning("Coop-avis for '%s' kunne ikke hentes: %s", retailer["name"], exc)
                selected[retailer["name"]] = None
            time.sleep(delay_seconds)
            continue

        url = f"https://etilbudsavis.dk/{quote(retailer['slug'], safe='')}"
        try:
            response = session.get(url, timeout=30)
            response.raise_for_status()
        except requests.RequestException as exc:
            LOGGER.warning("Aviser for '%s' kunne ikke hentes: %s", retailer["name"], exc)
            selected[retailer["name"]] = None
            continue
        successful_pages += 1
        publications = parse_publications(response.text)
        publication = select_week_publication(publications, week_number)
        selected[retailer["name"]] = publication
        if publication:
            LOGGER.info("%s: valgt '%s'", retailer["name"], publication.label)
        else:
            LOGGER.warning("%s: ingen tilbudsavis for uge %d er offentliggjort",
                           retailer["name"], week_number)
        time.sleep(delay_seconds)

    if not successful_pages:
        raise RuntimeError("Ingen butikssider kunne hentes fra eTilbudsavis.")
    return selected


def _filter_offers(offers: list[Offer], config: dict[str, Any],
                   publications: dict[str, Publication | None],
                   now: datetime | None = None,
                   uncategorized: list[Offer] | None = None) -> list[Offer]:
    now = now or datetime.now(timezone.utc)
    accepted: list[Offer] = []
    candidate_terms = _search_terms(config)
    for offer in offers:
        retailer_name = _retailer_display_name(offer.retailer, config["retailers"])
        publication = publications.get(retailer_name or "")
        searchable_text = f"{offer.name} {offer.description}"
        category = categorize_offer(
            offer.name, config["categories"], config.get("excluded_terms", []),
            offer.description,
        )
        if (not publication or offer.publication_id != publication.id
                or not _is_current(offer, now)):
            continue
        if not category:
            is_relevant_candidate = (
                publication.source != "coop"
                or any(_term_matches(searchable_text, term) for term in candidate_terms)
            )
            is_excluded = any(
                _term_matches(searchable_text, term)
                for term in config.get("excluded_terms", [])
            )
            if uncategorized is not None and is_relevant_candidate and not is_excluded:
                uncategorized.append(
                    Offer(
                        name=offer.name,
                        url=offer.url,
                        image_url=offer.image_url,
                        retailer=retailer_name,
                        price=offer.price,
                        valid_through=offer.valid_through,
                        publication_id=offer.publication_id,
                        description=offer.description,
                    )
                )
            continue
        accepted.append(
            Offer(
                name=offer.name,
                url=offer.url,
                image_url=offer.image_url,
                retailer=retailer_name,
                price=offer.price,
                valid_through=offer.valid_through,
                category=category,
                publication_id=offer.publication_id,
                description=offer.description,
            )
        )
    return accepted


def fetch_offers(config: dict[str, Any], publications: dict[str, Publication | None],
                 delay_seconds: float = 1.0,
                 uncategorized: list[Offer] | None = None,
                 week_number: int | None = None) -> list[Offer]:
    session = requests.Session()
    session.headers.update({"User-Agent": "OfferAutomation/1.0 (public offer search)"})
    unique_offers: dict[str, Offer] = {}
    successful_searches = 0

    for term in _search_terms(config):
        url = SEARCH_URL.format(quote(term, safe=""))
        try:
            response = session.get(url, timeout=30)
            response.raise_for_status()
        except requests.RequestException as exc:
            LOGGER.warning("Søgning efter '%s' fejlede: %s", term, exc)
            continue
        successful_searches += 1
        for offer in parse_search_results(response.text):
            unique_offers.setdefault(offer.url, offer)
        time.sleep(delay_seconds)

    for retailer in config["retailers"]:
        publication = publications.get(retailer["name"])
        if not publication or publication.source != "coop":
            continue
        try:
            coop_offers = _fetch_coop_offers(
                retailer, publication, week_number or date.today().isocalendar().week,
                session, delay_seconds,
            )
        except (requests.RequestException, ValueError) as exc:
            LOGGER.warning("Coop-tilbud for '%s' kunne ikke hentes: %s", retailer["name"], exc)
            continue
        for offer in coop_offers:
            unique_offers.setdefault(offer.url, offer)
        successful_searches += 1
        LOGGER.info("%s: hentede %d Coop-tilbud fra ugeavisen",
                    retailer["name"], len(coop_offers))

    if not successful_searches:
        raise RuntimeError("Ingen offentlige søgeresultater kunne hentes fra eTilbudsavis.")

    return _filter_offers(
        list(unique_offers.values()), config, publications, uncategorized=uncategorized
    )


def _add_text(slide, text: str, left: float, top: float, width: float, height: float,
              font_size: int, color: RGBColor, bold: bool = False,
              align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def _download_image(session: requests.Session, url: str) -> bytes:
    response = session.get(url, timeout=30)
    response.raise_for_status()
    return response.content


def _fit_image_dimensions(image_width: int, image_height: int,
                          max_width: float, max_height: float) -> tuple[float, float]:
    scale = min(max_width / image_width, max_height / image_height)
    return image_width * scale, image_height * scale


def _add_cover(prs: Presentation, retailer: str, week_number: int,
               publication: Publication | None, session: requests.Session) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 248, 245)
    _add_text(slide, retailer, 0.55, 0.2, 12.2, 0.55, 28, RGBColor(30, 57, 47), True)
    label = publication.label if publication else f"Ingen tilbudsavis for uge {week_number} fundet"
    _add_text(slide, label, 0.6, 0.78, 12.0, 0.4, 16, RGBColor(75, 85, 78))
    if not publication:
        return
    try:
        content = _download_image(session, publication.cover_image)
        with Image.open(BytesIO(content)) as image:
            width, height = _fit_image_dimensions(image.width, image.height, 5.0, 5.95)
        slide.shapes.add_picture(
            BytesIO(content),
            Inches((13.333 - width) / 2),
            Inches(1.35 + (5.95 - height) / 2),
            width=Inches(width),
            height=Inches(height),
        )
    except Exception as exc:
        LOGGER.warning("Forsiden for '%s' kunne ikke hentes: %s", retailer, exc)
        _add_text(slide, "Avisens forside kunne ikke hentes", 0.6, 3.0, 12.0, 0.6,
                  16, RGBColor(130, 130, 130), align=PP_ALIGN.CENTER)


def _add_category_slides(prs: Presentation, retailer: str, category: str,
                         offers: list[Offer], session: requests.Session) -> None:
    groups = [offers[index:index + 8] for index in range(0, len(offers), 8)] or [[]]
    for page_number, group in enumerate(groups, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        title = f"{category} | {retailer}"
        if len(groups) > 1:
            title += f" ({page_number})"
        _add_text(slide, title, 0.35, 0.2, 12.6, 0.55, 21, RGBColor(30, 57, 47), True)

        if not group:
            _add_text(slide, "Ingen avistilbud", 0.55, 2.6, 12.2, 0.8,
                      26, RGBColor(100, 105, 100), align=PP_ALIGN.CENTER)
            continue

        for index, offer in enumerate(group):
            column = index % 4
            row = index // 4
            left = 0.35 + column * 3.18
            top = 0.95 + row * 3.15
            image_left, image_top = left + 0.15, top
            image_width, image_height = 2.85, 2.48
            try:
                content = _download_image(session, offer.image_url)
                with Image.open(BytesIO(content)) as image:
                    width, height = _fit_image_dimensions(
                        image.width, image.height, image_width, image_height
                    )
                slide.shapes.add_picture(
                    BytesIO(content),
                    Inches(image_left + (image_width - width) / 2),
                    Inches(image_top + (image_height - height) / 2),
                    width=Inches(width),
                    height=Inches(height),
                )
            except Exception as exc:
                LOGGER.warning("Kunne ikke hente tilbudsbillede for '%s': %s", offer.name, exc)
                _add_text(slide, "Billede kunne ikke hentes", image_left, image_top,
                          image_width, image_height, 12, RGBColor(130, 130, 130),
                          align=PP_ALIGN.CENTER)
            caption = offer.name
            if offer.price:
                caption += f" | {offer.price} kr."
            _add_text(slide, caption[:90], left + 0.05, top + 2.5, 3.05, 0.55,
                      10, RGBColor(45, 50, 45), True, PP_ALIGN.CENTER)


def create_presentation(offers: list[Offer], retailers: list[dict[str, Any]],
                        publications: dict[str, Publication | None], week_number: int,
                        output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    session = requests.Session()

    for retailer_config in retailers:
        retailer = retailer_config["name"]
        _add_cover(prs, retailer, week_number, publications.get(retailer), session)
        for category in CATEGORY_ORDER:
            selected = [offer for offer in offers
                        if offer.retailer == retailer and offer.category == category]
            _add_category_slides(prs, retailer, category, selected, session)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def write_review_csv(offers: list[Offer], output_path: Path,
                     publications: dict[str, Publication | None]) -> None:
    import csv

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow(["Kategori", "Butik", "Avis-ID", "Tilbud", "Pris", "Udløber", "Link"])
        for offer in offers:
            writer.writerow([offer.category, offer.retailer,
                             publications[offer.retailer].id, offer.name, offer.price,
                             offer.valid_through, offer.url])


def write_uncategorized_csv(offers: list[Offer], output_path: Path) -> None:
    import csv

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow(["Butik", "Avis-ID", "Muligt tilbud", "Pris", "Udløber", "Link"])
        for offer in offers:
            writer.writerow([offer.retailer, offer.publication_id, offer.name,
                             offer.price, offer.valid_through, offer.url])


def main() -> None:
    parser = argparse.ArgumentParser(description="Lav ugens tilbudsavis-præsentation.")
    parser.add_argument("--config", type=Path, default=Path("config/categories.json"))
    parser.add_argument("--output", type=Path, default=Path("output/tilbudsoversigt.pptx"))
    parser.add_argument("--week", type=int, default=None,
                        help="ISO-ugenummer for tilbudsaviserne (standard: næste uge).")
    parser.add_argument("--delay", type=float, default=1.0,
                        help="Ventetid mellem offentlige søgninger i sekunder.")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    config = json.loads(args.config.read_text(encoding="utf-8"))
    today = date.today()
    next_monday = today + timedelta(days=7 - today.weekday())
    target_week = args.week or next_monday.isocalendar().week
    delay = max(0.0, args.delay)
    publications = fetch_publications(config["retailers"], target_week, delay_seconds=delay)
    uncategorized: list[Offer] = []
    offers = fetch_offers(config, publications, delay_seconds=delay,
                          uncategorized=uncategorized, week_number=target_week)
    create_presentation(offers, config["retailers"], publications, target_week, args.output)
    write_review_csv(offers, args.output.with_name("tilbud-til-gennemsyn.csv"), publications)
    write_uncategorized_csv(
        uncategorized, args.output.with_name("mulige-tilbud-til-kategorisering.csv")
    )
    LOGGER.info("Præsentation for uge %d oprettet med %d tilbud (%d til gennemsyn): %s",
                target_week, len(offers), len(uncategorized), args.output)


if __name__ == "__main__":
    main()
